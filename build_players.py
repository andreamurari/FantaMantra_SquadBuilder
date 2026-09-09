#!/usr/bin/env python3
"""
Rigenera la lista giocatori dentro index.html a partire dal file Excel del listone.

USO:
  1. Metti nella cartella il nuovo file Excel (nome tipo "Quotazioni_Fantacalcio_*.xlsx",
     scaricato da fantacalcio.it -> Quotazioni -> Esporta).
  2. Esegui:   python build_players.py
  3. Ricarica la pagina (o ripubblica l'artifact). Fatto.

Richiede:  pip install openpyxl
"""

import glob
import io
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Manca openpyxl. Installa con:  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
INDEX = os.path.join(HERE, "index.html")

# Il file xlsx di fantacalcio.it a volte ha caratteri accentati corrotti (�).
# Correzioni note per i nomi di Serie A.
FIX = {
    "Montip�": "Montipò", "Dod�": "Dodô", "Lucum�": "Lucumí",
    "Z� Pedro": "Zé Pedro", "Cand�": "Candé", "Dembel� A.": "Dembélé A.",
    "Kessi�": "Kessié", "Kon� M.": "Koné M.", "Cal�": "Calò",
    "Kon� I.": "Koné I.", "Bernab�": "Bernabé", "Ciss� A.": "Cissé A.",
    "Tour� I.": "Touré I.", "Traor� Hj.": "Traoré Hj.", "Soul�": "Soulé",
    "Laurient�": "Laurienté", "Tour� E.": "Touré E.",
}


def find_xlsx():
    cands = sorted(
        glob.glob(os.path.join(HERE, "*.xlsx")),
        key=os.path.getmtime, reverse=True,
    )
    cands = [c for c in cands if not os.path.basename(c).startswith("~$")]
    if not cands:
        sys.exit("Nessun file .xlsx trovato nella cartella.")
    return cands[0]


def main():
    xlsx = find_xlsx()
    print("Leggo:", os.path.basename(xlsx))
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Tutti"] if "Tutti" in wb.sheetnames else wb.worksheets[0]

    rows = list(ws.iter_rows(values_only=True))
    # riga 0 = titolo, riga 1 = intestazioni
    hdr = list(rows[1])
    idx = {name: i for i, name in enumerate(hdr)}
    need = ["Nome", "Squadra", "R", "RM", "Qt.A M", "FVM M"]
    for n in need:
        if n not in idx:
            sys.exit(f"Colonna '{n}' non trovata nel foglio. Intestazioni: {hdr}")

    out = []
    for r in rows[2:]:
        if not r or r[0] is None:
            continue
        name = str(r[idx["Nome"]])
        name = FIX.get(name, name).strip()
        team = r[idx["Squadra"]]
        macro = r[idx["R"]]
        rm = "/".join(x.strip() for x in str(r[idx["RM"]]).split(";"))
        qt = r[idx["Qt.A M"]]
        fvm = r[idx["FVM M"]]
        out.append([name, team, macro, rm, qt, fvm])

    left = [c[0] for c in out if "�" in c[0]]
    if left:
        print("ATTENZIONE, nomi ancora corrotti (aggiungili a FIX):", left)

    raw = "const RAW = " + json.dumps(out, ensure_ascii=False, separators=(",", ":")) + ";"

    src = io.open(INDEX, encoding="utf-8").read()
    new, n = re.subn(r"^const RAW = .*$", lambda m: raw, src, count=1, flags=re.M)
    if n != 1:
        sys.exit("Non ho trovato la riga 'const RAW = ...' in index.html")
    io.open(INDEX, "w", encoding="utf-8", newline="").write(new)

    print(f"OK: {len(out)} giocatori scritti in index.html")


if __name__ == "__main__":
    main()
